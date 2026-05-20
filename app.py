from flask import Flask, request, Response, render_template, stream_with_context, jsonify, redirect, url_for, session, flash
from flask_cors import CORS
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from functools import wraps
from pathlib import Path
from groq import Groq
from PyPDF2 import PdfReader
from docx import Document
from psycopg2.extras import RealDictCursor
import psycopg2
import json
import os
import re
import time

try:
    from dotenv import load_dotenv
    load_dotenv()
except Exception:
    pass

app = Flask(__name__)
app.secret_key = os.getenv("SECRET_KEY", "change-this-secret-key")
CORS(app, supports_credentials=True)

GROQ_API_KEY = os.getenv("GROQ_API_KEY")
MODEL_NAME = os.getenv("MODEL_NAME", "llama-3.1-8b-instant")
DATABASE_URL = os.getenv("DATABASE_URL")

client = Groq(api_key=GROQ_API_KEY) if GROQ_API_KEY else None

UPLOAD_FOLDER = Path("uploads")
UPLOAD_FOLDER.mkdir(exist_ok=True)

ALLOWED_EXTENSIONS = {"txt", "pdf", "docx"}
MAX_CONTEXT_CHARS = 6000


def get_db():
    if not DATABASE_URL:
        raise RuntimeError("DATABASE_URL is not configured")
    return psycopg2.connect(DATABASE_URL)


def init_db():
    sql = """
    CREATE TABLE IF NOT EXISTS users (
        id SERIAL PRIMARY KEY,
        username VARCHAR(100) UNIQUE NOT NULL,
        email VARCHAR(150) UNIQUE NOT NULL,
        password_hash TEXT NOT NULL,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    );

    CREATE TABLE IF NOT EXISTS chats (
        id SERIAL PRIMARY KEY,
        user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
        title VARCHAR(255) NOT NULL,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
        updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    );

    CREATE INDEX IF NOT EXISTS idx_chats_user_updated
    ON chats(user_id, updated_at);

    CREATE TABLE IF NOT EXISTS messages (
        id SERIAL PRIMARY KEY,
        chat_id INTEGER NOT NULL REFERENCES chats(id) ON DELETE CASCADE,
        user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
        role VARCHAR(20) NOT NULL CHECK (role IN ('user', 'assistant')),
        content TEXT NOT NULL,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    );

    CREATE INDEX IF NOT EXISTS idx_messages_chat_created
    ON messages(chat_id, created_at);

    CREATE INDEX IF NOT EXISTS idx_messages_user
    ON messages(user_id);

    CREATE TABLE IF NOT EXISTS uploaded_files (
        id SERIAL PRIMARY KEY,
        user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
        chat_id INTEGER REFERENCES chats(id) ON DELETE SET NULL,
        original_filename VARCHAR(255) NOT NULL,
        stored_filename VARCHAR(255) NOT NULL,
        file_path TEXT NOT NULL,
        extracted_text TEXT,
        uploaded_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    );

    CREATE INDEX IF NOT EXISTS idx_uploaded_files_user
    ON uploaded_files(user_id);

    CREATE INDEX IF NOT EXISTS idx_uploaded_files_chat
    ON uploaded_files(chat_id);
    """
    try:
        db = get_db()
        cur = db.cursor()
        cur.execute(sql)
        db.commit()
        cur.close()
        db.close()
        print("Database ready")
    except Exception as e:
        print("Database init failed:", e)


def login_required(fn):
    @wraps(fn)
    def wrapper(*args, **kwargs):
        if "user_id" not in session:
            if request.path.startswith("/api") or request.path in ["/chat", "/upload"]:
                return jsonify({"error": "Unauthorized"}), 401
            return redirect(url_for("login"))
        return fn(*args, **kwargs)
    return wrapper


def allowed_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS


def clean_text(text):
    text = re.sub(r"\s+", " ", text or "")
    return text.strip()


def read_file(path: Path):
    suffix = path.suffix.lower()
    text = ""

    if suffix == ".txt":
        text = path.read_text(encoding="utf-8", errors="ignore")

    elif suffix == ".pdf":
        reader = PdfReader(str(path))
        for page in reader.pages:
            text += (page.extract_text() or "") + "\n"

    elif suffix == ".docx":
        doc = Document(str(path))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        table_texts = []
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        table_texts.append(cell.text.strip())
        text = "\n".join(paragraphs + table_texts)

    return clean_text(text)


def user_owns_chat(chat_id, user_id):
    db = get_db()
    cur = db.cursor(cursor_factory=RealDictCursor)
    cur.execute("SELECT id FROM chats WHERE id=%s AND user_id=%s", (chat_id, user_id))
    row = cur.fetchone()
    cur.close()
    db.close()
    return row is not None


def create_chat_db(user_id, title):
    title = (title or "New Chat").strip()[:255]
    db = get_db()
    cur = db.cursor(cursor_factory=RealDictCursor)
    cur.execute(
        "INSERT INTO chats (user_id, title) VALUES (%s, %s) RETURNING id",
        (user_id, title)
    )
    row = cur.fetchone()
    db.commit()
    cur.close()
    db.close()
    return row["id"]


def save_message_db(chat_id, user_id, role, content):
    db = get_db()
    cur = db.cursor()
    cur.execute(
        "INSERT INTO messages (chat_id, user_id, role, content) VALUES (%s, %s, %s, %s)",
        (chat_id, user_id, role, content)
    )
    cur.execute(
        "UPDATE chats SET updated_at=CURRENT_TIMESTAMP WHERE id=%s AND user_id=%s",
        (chat_id, user_id)
    )
    db.commit()
    cur.close()
    db.close()


def get_chat_history_db(chat_id, user_id, limit=8):
    db = get_db()
    cur = db.cursor(cursor_factory=RealDictCursor)
    cur.execute(
        """
        SELECT m.role, m.content
        FROM messages m
        JOIN chats c ON m.chat_id = c.id
        WHERE m.chat_id=%s AND c.user_id=%s AND m.user_id=%s
        ORDER BY m.created_at DESC, m.id DESC
        LIMIT %s
        """,
        (chat_id, user_id, user_id, limit)
    )
    rows = cur.fetchall()[::-1]
    cur.close()
    db.close()
    return [{"role": r["role"], "content": r["content"]} for r in rows]


def save_uploaded_file_db(user_id, chat_id, original_filename, stored_filename, file_path, extracted_text):
    db = get_db()
    cur = db.cursor()
    cur.execute(
        """
        INSERT INTO uploaded_files
        (user_id, chat_id, original_filename, stored_filename, file_path, extracted_text)
        VALUES (%s, %s, %s, %s, %s, %s)
        """,
        (user_id, chat_id, original_filename, stored_filename, file_path, extracted_text)
    )
    db.commit()
    cur.close()
    db.close()


def get_user_documents(user_id, chat_id=None, limit=5):
    db = get_db()
    cur = db.cursor(cursor_factory=RealDictCursor)

    if chat_id:
        cur.execute(
            """
            SELECT original_filename, extracted_text, uploaded_at
            FROM uploaded_files
            WHERE user_id=%s AND (chat_id=%s OR chat_id IS NULL)
            ORDER BY uploaded_at DESC
            LIMIT %s
            """,
            (user_id, chat_id, limit)
        )
    else:
        cur.execute(
            """
            SELECT original_filename, extracted_text, uploaded_at
            FROM uploaded_files
            WHERE user_id=%s
            ORDER BY uploaded_at DESC
            LIMIT %s
            """,
            (user_id, limit)
        )

    rows = cur.fetchall()
    cur.close()
    db.close()
    return rows


def make_relevant_context(user_message, user_id, chat_id=None):
    docs = get_user_documents(user_id, chat_id, limit=5)

    if not docs:
        return "", []

    query_words = set(re.findall(r"[a-zA-Z0-9]+", user_message.lower()))
    selected_parts = []
    sources = []

    for doc in docs:
        text = doc.get("extracted_text") or ""
        filename = doc.get("original_filename")

        if not text:
            continue

        sentences = re.split(r"(?<=[.!?])\s+", text)
        scored = []

        for sentence in sentences:
            sentence_clean = sentence.strip()
            if len(sentence_clean) < 30:
                continue

            words = set(re.findall(r"[a-zA-Z0-9]+", sentence_clean.lower()))
            score = len(query_words.intersection(words))

            if score > 0:
                scored.append((score, sentence_clean))

        scored.sort(key=lambda x: x[0], reverse=True)

        if scored:
            best_text = " ".join([s for _, s in scored[:20]])
        else:
            lower_q = user_message.lower()
            if any(w in lower_q for w in ["summarize", "summary", "explain document", "uploaded file", "file", "document"]):
                best_text = text[:2500]
            else:
                best_text = ""

        if best_text:
            selected_parts.append(f"[Document: {filename}]\n{best_text[:3000]}")
            sources.append(filename)

    context = "\n\n---\n\n".join(selected_parts)
    context = context[:MAX_CONTEXT_CHARS]

    unique_sources = []
    seen = set()
    for src in sources:
        if src not in seen:
            unique_sources.append(src)
            seen.add(src)

    return context, unique_sources


def build_ai_messages(user_message, chat_id, user_id):
    history = get_chat_history_db(chat_id, user_id, limit=8)
    context, sources = make_relevant_context(user_message, user_id, chat_id)

    system_prompt = (
        "You are DeepKul-AI, a helpful and professional AI assistant. "
        "If document context is provided and relevant, use it to answer. "
        "If document context is not relevant or not provided, answer normally using your own knowledge. "
        "Give clear, accurate, structured answers. "
        "Do not invent document facts. "
        "Do not mention sources inside the answer; sources will be appended separately."
    )

    messages = [{"role": "system", "content": system_prompt}]
    messages.extend(history)

    if context:
        user_content = f"DOCUMENT CONTEXT:\n{context}\n\nUSER QUESTION:\n{user_message}"
    else:
        user_content = user_message

    messages.append({"role": "user", "content": user_content})
    return messages, sources


@app.route("/register", methods=["GET", "POST"])
def register():
    if request.method == "GET":
        return render_template("register.html")

    username = request.form.get("username", "").strip()
    email = request.form.get("email", "").strip().lower()
    password = request.form.get("password", "")

    if not username or not email or not password:
        flash("All fields are required.", "error")
        return redirect(url_for("register"))

    try:
        db = get_db()
        cur = db.cursor()
        cur.execute(
            "INSERT INTO users (username, email, password_hash) VALUES (%s, %s, %s)",
            (username, email, generate_password_hash(password))
        )
        db.commit()
        cur.close()
        db.close()
        flash("Account created successfully. Please login.", "success")
        return redirect(url_for("login"))

    except Exception as e:
        flash("Username or email already exists.", "error")
        return redirect(url_for("register"))


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "GET":
        return render_template("login.html")

    identifier = request.form.get("identifier", "").strip()
    password = request.form.get("password", "")

    db = get_db()
    cur = db.cursor(cursor_factory=RealDictCursor)
    cur.execute(
        "SELECT * FROM users WHERE username=%s OR email=%s",
        (identifier, identifier)
    )
    user = cur.fetchone()
    cur.close()
    db.close()

    if user and check_password_hash(user["password_hash"], password):
        session["user_id"] = user["id"]
        session["username"] = user["username"]
        return redirect(url_for("home"))

    flash("Invalid login details.", "error")
    return redirect(url_for("login"))


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/")
@login_required
def home():
    return render_template("index.html", username=session.get("username"))


@app.route("/upload", methods=["POST"])
@login_required
def upload():
    user_id = session["user_id"]
    chat_id = request.form.get("chat_id")

    if chat_id:
        chat_id = int(chat_id)
        if not user_owns_chat(chat_id, user_id):
            return jsonify({"error": "Unauthorized chat access"}), 403
    else:
        chat_id = None

    files = request.files.getlist("files")
    if not files and request.files.get("file"):
        files = [request.files.get("file")]

    if not files:
        return jsonify({"error": "No files uploaded"}), 400

    user_upload_dir = UPLOAD_FOLDER / str(user_id)
    user_upload_dir.mkdir(parents=True, exist_ok=True)

    uploaded = []
    errors = []

    for file in files:
        if not file or file.filename == "":
            continue

        if not allowed_file(file.filename):
            errors.append(f"{file.filename}: unsupported file type")
            continue

        original_filename = file.filename
        safe_name = secure_filename(original_filename)
        stored_filename = f"{int(time.time())}_{safe_name}"
        path = user_upload_dir / stored_filename

        try:
            file.save(path)
            extracted_text = read_file(path)

            save_uploaded_file_db(
                user_id=user_id,
                chat_id=chat_id,
                original_filename=original_filename,
                stored_filename=stored_filename,
                file_path=str(path),
                extracted_text=extracted_text
            )

            uploaded.append({
                "filename": original_filename,
                "status": "ok"
            })

        except Exception as e:
            errors.append(f"{original_filename}: {str(e)}")

    return jsonify({
        "message": f"Uploaded {len(uploaded)} file(s)",
        "files": uploaded,
        "errors": errors
    })


@app.route("/chat", methods=["POST"])
@login_required
def chat():
    data = request.get_json(silent=True) or {}
    user_message = data.get("message", "").strip()
    chat_id = data.get("chat_id")

    if not user_message:
        return Response("No message provided", status=400)

    user_id = session["user_id"]

    if not chat_id:
        chat_id = create_chat_db(user_id, user_message[:60])
    else:
        chat_id = int(chat_id)
        if not user_owns_chat(chat_id, user_id):
            return jsonify({"error": "Unauthorized chat access"}), 403

    save_message_db(chat_id, user_id, "user", user_message)

    messages, sources = build_ai_messages(user_message, chat_id, user_id)

    def generate():
        full_response = ""

        try:
            if client is None:
                raise RuntimeError("GROQ_API_KEY is not configured")

            yield f"data: {json.dumps({'chat_id': chat_id})}\n\n"

            response = client.chat.completions.create(
                model=MODEL_NAME,
                messages=messages,
                temperature=0.45,
                max_tokens=2048,
                stream=True
            )

            for chunk in response:
                content = chunk.choices[0].delta.content
                if content:
                    full_response += content
                    yield f"data: {json.dumps({'token': content})}\n\n"

            if sources:
                source_text = "\n\n**Sources:**\n" + "\n".join([f"- {src}" for src in sources])
                full_response += source_text
                yield f"data: {json.dumps({'token': source_text})}\n\n"

            save_message_db(chat_id, user_id, "assistant", full_response)

            yield f"data: {json.dumps({'done': True, 'chat_id': chat_id})}\n\n"

        except Exception as e:
            yield f"data: {json.dumps({'error': str(e)})}\n\n"

    return Response(
        stream_with_context(generate()),
        mimetype="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "X-Accel-Buffering": "no"
        }
    )


@app.route("/api/chats", methods=["GET"])
@login_required
def api_get_chats():
    db = get_db()
    cur = db.cursor(cursor_factory=RealDictCursor)
    cur.execute(
        """
        SELECT id, title, created_at, updated_at
        FROM chats
        WHERE user_id=%s
        ORDER BY updated_at DESC, id DESC
        """,
        (session["user_id"],)
    )
    rows = cur.fetchall()
    cur.close()
    db.close()

    for row in rows:
        row["created_at"] = str(row["created_at"])
        row["updated_at"] = str(row["updated_at"])

    return jsonify({"chats": rows})


@app.route("/api/chats", methods=["POST"])
@login_required
def api_create_chat():
    data = request.get_json(silent=True) or {}
    title = data.get("title") or "New Chat"
    chat_id = create_chat_db(session["user_id"], title)
    return jsonify({"id": chat_id, "title": title})


@app.route("/api/chats/<int:chat_id>/messages", methods=["GET"])
@login_required
def api_get_messages(chat_id):
    if not user_owns_chat(chat_id, session["user_id"]):
        return jsonify({"error": "Unauthorized"}), 403

    db = get_db()
    cur = db.cursor(cursor_factory=RealDictCursor)
    cur.execute(
        """
        SELECT role, content, created_at
        FROM messages
        WHERE chat_id=%s AND user_id=%s
        ORDER BY created_at ASC, id ASC
        """,
        (chat_id, session["user_id"])
    )
    rows = cur.fetchall()
    cur.close()
    db.close()

    messages = [
        {
            "role": r["role"],
            "content": r["content"],
            "created_at": str(r["created_at"])
        }
        for r in rows
    ]

    return jsonify({"messages": messages})


@app.route("/api/chats/<int:chat_id>", methods=["PATCH"])
@login_required
def api_rename_chat(chat_id):
    if not user_owns_chat(chat_id, session["user_id"]):
        return jsonify({"error": "Unauthorized"}), 403

    data = request.get_json(silent=True) or {}
    title = (data.get("title") or "New Chat").strip()[:255]

    db = get_db()
    cur = db.cursor()
    cur.execute(
        "UPDATE chats SET title=%s, updated_at=CURRENT_TIMESTAMP WHERE id=%s AND user_id=%s",
        (title, chat_id, session["user_id"])
    )
    db.commit()
    cur.close()
    db.close()

    return jsonify({"success": True})


@app.route("/api/chats/<int:chat_id>", methods=["DELETE"])
@login_required
def api_delete_chat(chat_id):
    if not user_owns_chat(chat_id, session["user_id"]):
        return jsonify({"error": "Unauthorized"}), 403

    db = get_db()
    cur = db.cursor()
    cur.execute(
        "DELETE FROM chats WHERE id=%s AND user_id=%s",
        (chat_id, session["user_id"])
    )
    db.commit()
    cur.close()
    db.close()

    return jsonify({"success": True})


@app.route("/health")
def health():
    try:
        db = get_db()
        cur = db.cursor()
        cur.execute("SELECT 1")
        cur.close()
        db.close()
        db_status = "connected"
    except Exception:
        db_status = "not connected"

    return jsonify({
        "status": "ok",
        "mode": "lightweight-document-context",
        "database": db_status,
        "model": MODEL_NAME
    })


init_db()

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port, threaded=True, debug=False)